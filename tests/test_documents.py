"""Document extraction tests. Files are built in code, so no fixtures and no LLM."""
import pytest

import documents
from documents import DocumentError, Passage, Document, extract, select_passages


def make_pdf(path, pages):
    """A tiny valid PDF with one line of text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
                    f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>")
        stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(out)


def test_text_and_markdown_split_into_paragraphs(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("Acme has 500 employees.\n\nIt was founded in 2001.")
    doc = extract(f)
    assert [p.label for p in doc.passages] == ["notes.txt para 1", "notes.txt para 2"]
    assert "founded in 2001" in doc.text


def test_csv_rows_are_labelled(tmp_path):
    f = tmp_path / "data.csv"
    f.write_text("org,employees\nacme,500\n")
    assert [p.label for p in extract(f).passages] == ["data.csv row 1", "data.csv row 2"]


def test_pdf_pages_are_labelled(tmp_path):
    f = tmp_path / "report.pdf"
    make_pdf(f, ["Acme has 500 employees", "Revenue grew in 2025"])
    doc = extract(f)
    assert [p.label for p in doc.passages] == ["report.pdf p.1", "report.pdf p.2"]
    assert "500 employees" in doc.passages[0].text


def test_pdf_without_text_is_reported_as_scanned(tmp_path):
    f = tmp_path / "scan.pdf"
    make_pdf(f, [""])
    with pytest.raises(DocumentError, match="no readable text"):
        extract(f)


def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document as Docx
    d = Docx()
    d.add_paragraph("Acme has 500 employees.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "HQ", "Berlin"
    f = tmp_path / "memo.docx"
    d.save(str(f))
    doc = extract(f)
    assert "Acme has 500 employees." in doc.text
    assert any(p.label == "memo.docx table 1 row 1" and p.text == "HQ | Berlin" for p in doc.passages)


def test_xlsx_rows_carry_sheet_name_and_row_number(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append(["org", "employees"])
    ws.append(["acme", 500])
    f = tmp_path / "staff.xlsx"
    wb.save(str(f))
    doc = extract(f)
    assert [(p.label, p.text) for p in doc.passages] == [
        ("staff.xlsx Staff row 1", "org | employees"), ("staff.xlsx Staff row 2", "acme | 500")]


def test_old_doc_without_libreoffice_gives_clear_error(tmp_path, monkeypatch):
    monkeypatch.setattr(documents.shutil, "which", lambda name: None)
    f = tmp_path / "old.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocumentError, match="save it as .docx"):
        extract(f)


def test_unsupported_missing_corrupt_and_oversized_files(tmp_path, monkeypatch):
    img = tmp_path / "photo.png"
    img.write_bytes(b"x")
    with pytest.raises(DocumentError, match="unsupported"):
        extract(img)
    with pytest.raises(DocumentError, match="not found"):
        extract(tmp_path / "missing.txt")
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(DocumentError, match="could not read"):
        extract(bad)
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 3)
    big = tmp_path / "big.txt"
    big.write_text("too long")
    with pytest.raises(DocumentError, match="larger than"):
        extract(big)


def test_short_documents_are_sent_whole():
    doc = Document("a.txt", [Passage("a p1", "one"), Passage("a p2", "two")])
    assert select_passages(doc, "anything") == doc.passages


def test_long_documents_keep_the_relevant_passages_in_order():
    filler = [Passage(f"d p{i}", "lorem ipsum dolor " * 20) for i in range(10)]
    hit1 = Passage("d hit1", "Acme employees total 500 people")
    hit2 = Passage("d hit2", "the number of Acme employees is stable")
    doc = Document("d", filler[:3] + [hit2] + filler[3:] + [hit1])
    chosen = select_passages(doc, "Acme has 500 employees", max_chars=200)
    assert [p.label for p in chosen] == ["d hit2", "d hit1"]   # document order, filler dropped


def test_long_document_with_nothing_relevant_sends_nothing():
    doc = Document("d", [Passage(f"d p{i}", "lorem ipsum dolor " * 20) for i in range(10)])
    assert select_passages(doc, "Acme has 500 employees", max_chars=200) == []
