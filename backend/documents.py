"""Turn an uploaded file into labelled plain-text passages, for checking a fact against it.

extract(path) -> Document(name, passages). Each passage has a label saying where the text
came from ("report.pdf p.3", "data.xlsx Sheet1 row 4"), so a recorded source can be traced.
select_passages(doc, claim) keeps what is sent to the LLM small: if the whole document fits
it is sent as is, otherwise only the passages sharing the most words with the claim.

No LLM and no database here. Supported: .txt .md .csv .pdf (text PDFs) .docx .xlsx, and
.doc when LibreOffice (soffice) is installed.
"""
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_CHARS = 30_000  # most text sent to the checker for one document

_STOPWORDS = {
    "the", "and", "for", "are", "was", "were", "has", "have", "had", "with", "that", "this",
    "from", "its", "their", "than", "then", "into", "not", "but", "all", "any", "can", "per",
}


class DocumentError(Exception):
    """The file can't be read as a document (unsupported type, too big, scanned, corrupt)."""


@dataclass
class Passage:
    label: str  # where it came from, e.g. "report.pdf p.3"
    text: str


@dataclass
class Document:
    name: str
    passages: list

    @property
    def text(self):
        return "\n".join(p.text for p in self.passages)


def extract(path):
    path = Path(path)
    if not path.is_file():
        raise DocumentError(f"file not found: {path}")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise DocumentError(f"{path.name} is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise DocumentError(
            f"unsupported file type {path.suffix!r}; use one of {sorted(_READERS)}")
    try:
        passages = reader(path)
    except DocumentError:
        raise
    except Exception as e:  # corrupt or password-protected files raise many different errors
        raise DocumentError(f"could not read {path.name}: {e}") from e
    passages = [p for p in passages if p.text.strip()]
    if not passages:
        raise DocumentError(f"{path.name} has no readable text (a scanned file needs OCR)")
    return Document(path.name, passages)


def select_passages(doc, claim, max_chars=MAX_CHARS):
    """Passages to send to the checker, in document order, at most max_chars of text."""
    if len(doc.text) <= max_chars:
        return list(doc.passages)
    claim_words = _words(claim)
    scored = [(len(claim_words & _words(p.text)), -i, p) for i, p in enumerate(doc.passages)]
    chosen, used = [], 0
    for score, neg_i, p in sorted(scored, key=lambda s: (s[0], s[1]), reverse=True):
        if score == 0:
            break
        if used + len(p.text) > max_chars:
            continue
        chosen.append((-neg_i, p))
        used += len(p.text)
    return [p for _, p in sorted(chosen, key=lambda c: c[0])]


def _words(text):
    return {w for w in re.findall(r"[a-z0-9]{3,}", text.lower()) if w not in _STOPWORDS}


# Readers: one per file type, each returns a list of Passage

def _read_text(path):
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".csv":
        return [Passage(f"{path.name} row {i}", line)
                for i, line in enumerate(text.splitlines(), start=1)]
    paras = [p.strip() for p in re.split(r"\n\s*\n", text)]
    return [Passage(f"{path.name} para {i}", p) for i, p in enumerate(paras, start=1)]


def _read_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentError(f"{path.name} is password-protected")
    return [Passage(f"{path.name} p.{i}", page.extract_text() or "")
            for i, page in enumerate(reader.pages, start=1)]


def _read_docx(path):
    from docx import Document as Docx
    d = Docx(str(path))
    out = [Passage(f"{path.name} para {i}", p.text) for i, p in enumerate(d.paragraphs, start=1)]
    for t, table in enumerate(d.tables, start=1):
        for r, row in enumerate(table.rows, start=1):
            out.append(Passage(f"{path.name} table {t} row {r}",
                               " | ".join(c.text.strip() for c in row.cells)))
    return out


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append(Passage(f"{path.name} {ws.title} row {r}", " | ".join(cells)))
    wb.close()
    return out


def _read_doc(path):
    """Old binary Word: needs LibreOffice to convert it to .docx first."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise DocumentError(
            f"{path.name} is an old .doc file; save it as .docx (or install LibreOffice)")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(path)],
                       check=True, capture_output=True, timeout=120)
        converted = Path(tmp) / (path.stem + ".docx")
        return [Passage(p.label.replace(converted.name, path.name), p.text)
                for p in _read_docx(converted)]


_READERS = {
    ".txt": _read_text, ".md": _read_text, ".csv": _read_text,
    ".pdf": _read_pdf, ".docx": _read_docx, ".xlsx": _read_xlsx, ".doc": _read_doc,
}
